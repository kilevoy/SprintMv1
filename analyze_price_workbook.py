from __future__ import annotations

import json
import math
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"E:\SprintMv1\input_price_review\Прайс для предрасчетов (не изменять) — 23092026.xlsx")
OUT = Path(r"E:\SprintMv1\input_price_review\price_analysis.json")

PRICE_WORDS = ("цена", "стоим", "прайс", "руб", "р/", "руб.")
NAME_WORDS = ("наимен", "номен", "материал", "профиль", "изделие", "позиция")
UNIT_WORDS = ("ед.изм", "ед. изм", "единиц", "ед изм")


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\n", " ")).strip()


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def percentile(values, p):
    if not values:
        return None
    xs = sorted(values)
    k = (len(xs) - 1) * p
    lo, hi = math.floor(k), math.ceil(k)
    if lo == hi:
        return xs[lo]
    return xs[lo] * (hi - k) + xs[hi] * (k - lo)


def header_text(ws, row, col):
    parts = []
    for rr in range(max(1, row - 2), row + 1):
        s = clean(ws.cell(rr, col).value)
        if s and not s.startswith("="):
            parts.append(s)
    return " | ".join(parts)


def nearest_name(ws, row, price_col):
    candidates = []
    for c in range(max(1, price_col - 6), price_col):
        v = clean(ws.cell(row, c).value)
        if v and not v.startswith("=") and not re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", v):
            candidates.append((c, v))
    return candidates[-1][1] if candidates else ""


def main():
    # Normal mode is intentional here: read-only worksheets make repeated random
    # access to neighbouring name/price cells prohibitively slow on this workbook.
    wb_formula = load_workbook(SOURCE, read_only=False, data_only=False, keep_links=True)
    wb_values = load_workbook(SOURCE, read_only=False, data_only=True, keep_links=True)
    report = {
        "file": str(SOURCE),
        "sheets": [],
        "price_columns": [],
        "duplicate_names": [],
        "external_links": [],
    }

    duplicates = defaultdict(list)
    external_refs = Counter()

    for name in wb_formula.sheetnames:
        wsf = wb_formula[name]
        wsv = wb_values[name]
        sheet = {
            "name": name,
            "max_row": wsf.max_row,
            "max_column": wsf.max_column,
            "nonempty": 0,
            "formulas": 0,
            "formula_errors_cached": 0,
            "price_headers": [],
        }

        candidate_headers = []
        scan_rows = min(wsf.max_row, 80)
        scan_cols = min(wsf.max_column, 80)
        for row in wsf.iter_rows(min_row=1, max_row=scan_rows, min_col=1, max_col=scan_cols):
            for cell in row:
                v = cell.value
                if v not in (None, ""):
                    sheet["nonempty"] += 1
                if isinstance(v, str) and v.startswith("="):
                    sheet["formulas"] += 1
                    for ref in re.findall(r"\[[^\]]+\][^!]+!", v):
                        external_refs[ref] += 1
                txt = clean(v).lower()
                if txt and any(w in txt for w in PRICE_WORDS):
                    candidate_headers.append((cell.row, cell.column, clean(v)))

        # Count formulas and external refs beyond the header scan.
        for row in wsf.iter_rows(min_row=scan_rows + 1):
            for cell in row:
                v = cell.value
                if v not in (None, ""):
                    sheet["nonempty"] += 1
                if isinstance(v, str) and v.startswith("="):
                    sheet["formulas"] += 1
                    for ref in re.findall(r"\[[^\]]+\][^!]+!", v):
                        external_refs[ref] += 1

        seen_cols = set()
        for hr, pc, label in candidate_headers:
            # Avoid treating paragraphs/titles as price columns: require numeric content below.
            vals = []
            zero_cells = []
            negative_cells = []
            formula_cells = 0
            blank_after_items = []
            entries = []
            for r in range(hr + 1, wsf.max_row + 1):
                fv = wsf.cell(r, pc).value
                vv = wsv.cell(r, pc).value
                if isinstance(fv, str) and fv.startswith("="):
                    formula_cells += 1
                if num(vv):
                    vals.append(float(vv))
                    if vv == 0:
                        zero_cells.append(f"{wsf.cell(r, pc).coordinate}")
                    elif vv < 0:
                        negative_cells.append(f"{wsf.cell(r, pc).coordinate}")
                    item = nearest_name(wsv, r, pc)
                    if item:
                        entries.append((r, item, float(vv), wsf.cell(r, pc).coordinate))
                        norm = re.sub(r"[^a-zа-яё0-9]+", " ", item.lower()).strip()
                        if norm:
                            duplicates[norm].append((name, wsf.cell(r, pc).coordinate, item, float(vv)))
                else:
                    item = nearest_name(wsv, r, pc)
                    if item and not isinstance(fv, str) and r <= hr + 500:
                        blank_after_items.append((wsf.cell(r, pc).coordinate, item))

            positive = [x for x in vals if x > 0]
            if len(vals) < 3 or pc in seen_cols:
                continue
            seen_cols.add(pc)
            q1 = percentile(positive, 0.25)
            q3 = percentile(positive, 0.75)
            med = percentile(positive, 0.5)
            outliers = []
            if q1 is not None and q3 is not None:
                iqr = q3 - q1
                lo = max(0, q1 - 3 * iqr)
                hi = q3 + 3 * iqr
                for r, item, value, coord in entries:
                    if value > 0 and (value < lo or value > hi):
                        outliers.append({"cell": coord, "item": item, "value": value})
            col = {
                "sheet": name,
                "header_cell": wsf.cell(hr, pc).coordinate,
                "header": label,
                "context_header": header_text(wsv, hr, pc),
                "numeric_count": len(vals),
                "formula_count": formula_cells,
                "positive_count": len(positive),
                "zero_count": len(zero_cells),
                "negative_count": len(negative_cells),
                "min_positive": min(positive) if positive else None,
                "median_positive": med,
                "max_positive": max(positive) if positive else None,
                "sum": sum(vals),
                "zero_cells_sample": zero_cells[:30],
                "negative_cells_sample": negative_cells[:30],
                "blank_price_with_item_sample": [
                    {"cell": c, "item": i} for c, i in blank_after_items[:30]
                ],
                "outliers_sample": sorted(outliers, key=lambda x: x["value"], reverse=True)[:30],
            }
            sheet["price_headers"].append(wsf.cell(hr, pc).coordinate)
            report["price_columns"].append(col)

        report["sheets"].append(sheet)

    for norm, rows in duplicates.items():
        prices = sorted({round(x[3], 8) for x in rows})
        if len(rows) > 1 and len(prices) > 1:
            report["duplicate_names"].append({
                "normalized_name": norm,
                "distinct_prices": prices,
                "occurrences": [
                    {"sheet": s, "cell": c, "item": i, "price": p} for s, c, i, p in rows
                ],
            })
    report["duplicate_names"].sort(key=lambda x: (-len(x["occurrences"]), x["normalized_name"]))
    report["external_links"] = [
        {"reference": k, "formula_count": v} for k, v in external_refs.most_common()
    ]
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "sheet_count": len(report["sheets"]),
        "price_column_count": len(report["price_columns"]),
        "duplicate_conflict_count": len(report["duplicate_names"]),
        "external_reference_count": sum(x["formula_count"] for x in report["external_links"]),
        "output": str(OUT),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
