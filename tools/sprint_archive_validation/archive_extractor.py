from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

from .archive_classifier import classify
from .config import INPUT_CELLS, OUTPUT_CELLS, BRANCH_OUTPUT_CELLS, CLIMATE_ROW_CELLS
from .profile_normalizer import parse_profile


REQUIRED_INPUTS = {"city", "span_m", "length_m", "height_m", "responsibility_level", "roof_covering", "deck_grade"}


def _clean(value: Any):
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return None if stripped == "" else stripped
    return value


def _numeric(value: Any):
    value = _clean(value)
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return value


def _is_error(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("#")


def _cell(wb, sheet_name: str, address: str):
    if sheet_name not in wb.sheetnames:
        return None, None
    formula_wb = wb
    return _clean(formula_wb[sheet_name][address].value), None


def _find_city_row(wb, city: str):
    if not city or "снегветер" not in wb.sheetnames:
        return None
    sheet = wb["снегветер"]
    for row in range(3, 601):
        if _clean(sheet.cell(row, 2).value) == city:
            return row
    return None


def extract_archive(scan_record: dict) -> dict:
    path = Path(scan_record["source_file"])
    formula_wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    value_wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    record = dict(scan_record)
    record["source_modified_date"] = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat()
    record["quality_flags"] = list(record.get("quality_flags", []))
    record["provenance"] = {}
    record["inputs"] = {}
    record["archive"] = {}
    record["country"] = None
    record["normative_system"] = "SP_20"
    record["provenance"]["country_source"] = None
    record["provenance"]["normative_system_source"] = "снегветер!F:I canonical SP 20 path"
    record["search_text"] = [path.name, path.parent.name]
    if record.get("archive_title"):
        record["search_text"].append(record["archive_title"])

    if "вывод" not in value_wb.sheetnames:
        record["archive_layout"] = "UNSUPPORTED_RESULT_WORKBOOK"
        record["archive_sheetnames"] = list(value_wb.sheetnames)
        record["parse_error"] = "UNSUPPORTED_ARCHIVE_LAYOUT"
        record["quality_flags"].extend(["UNSUPPORTED_ARCHIVE_LAYOUT", "INPUTS_NOT_EXTRACTED"])
        record["provenance"]["archive_layout_source"] = "workbook sheet inventory"
        record["provenance"]["input_cells_available"] = False
        formula_wb.close()
        value_wb.close()
        return classify(record)

    record["archive_layout"] = "STANDARD_SPRINT_SOURCE_SELECTION"
    record["archive_sheetnames"] = list(value_wb.sheetnames)

    for field, address in INPUT_CELLS.items():
        value = _clean(value_wb["вывод"][address].value)
        if field in {"span_m", "length_m", "height_m", "responsibility_level", "frame_step_input_m", "gate_le_6_count", "gate_gt_6_count", "door_count", "window_height_m", "strip_window_length_m", "separate_window_count"}:
            value = _numeric(value)
        record["inputs"][field] = value
        record["provenance"][f"input_{field}_source"] = f"вывод!{address}"
        if value is None and field in REQUIRED_INPUTS:
            record["quality_flags"].append("MISSING_CITY" if field == "city" else "INPUT_INCOMPLETE")
        if value is None and field == "deck_grade":
            record["quality_flags"].append("MISSING_DECK_GRADE")
    city = record["inputs"].get("city")
    record["search_text"].extend(str(value) for value in record["inputs"].values() if value is not None)

    output_values = {}
    for field, address in OUTPUT_CELLS.items():
        value = _clean(value_wb["вывод"][address].value)
        output_values[field] = _numeric(value) if field.endswith(("_m", "_mm", "_kg", "_kg_m2", "_kpa", "_t")) else value
        record["provenance"][f"archive_{field}_source"] = f"вывод!{address}"

    e8 = _numeric(value_wb["вывод"]["E8"].value)
    e9 = _numeric(value_wb["вывод"]["E9"].value)
    active_branch = 15 if isinstance(e8, (int, float)) and isinstance(e9, (int, float)) and e8 > e9 else 14
    family_row = {9: 2, 12: 3, 15: 4, 18: 5, 21: 6, 24: 7}.get(record["inputs"].get("span_m"))
    branch_sheet = value_wb["подбор"] if "подбор" in value_wb.sheetnames else None
    frame_mass = _numeric(branch_sheet[f"G{family_row}"].value) if branch_sheet and family_row else None
    secondary_mass_m2 = _numeric(branch_sheet[f"H{family_row}"].value) if branch_sheet and family_row else None
    record["archive"] = {
        "snow_region": output_values.get("snow_region_display"),
        "snow_load_kpa": None,
        "wind_region": output_values.get("wind_region_display"),
        "wind_load_kpa": None,
        "frame_step_m": output_values.get("frame_step_m"),
        "frame_count": None,
        "beam_profile": output_values.get("beam_profile"),
        "beam_steel": output_values.get("beam_steel"),
        "column_profile": output_values.get("column_profile"),
        "column_steel": output_values.get("column_steel"),
        "frame_mass_kg": frame_mass,
        "purlin_profile": output_values.get("purlin_profile"),
        "purlin_steel": output_values.get("purlin_steel"),
        "purlin_step_mm": output_values.get("purlin_step_mm"),
        "purlin_mass_kg": _numeric(value_wb["вывод"]["E24"].value),
        "secondary_mass_kg": None,
        "secondary_mass_kg_m2": secondary_mass_m2,
        "opening_mass_kg_m2": output_values.get("opening_mass_kg_m2"),
        "structural_base_kg_m2": output_values.get("structural_base_kg_m2"),
        "d69_kg_m2": output_values.get("d69_kg_m2"),
        "frame_total_kg": None,
        "active_branch": active_branch,
        "alternate_structural_base_kg_m2": output_values.get("alternate_structural_base_kg_m2"),
        "canonical_snow_region": None,
        "canonical_wind_region": None,
        "canonical_snow_load_kpa": None,
        "canonical_wind_load_kpa": None,
    }
    record["provenance"].update({
        "archive_frame_mass_kg_source": f"подбор!G{family_row}" if family_row else None,
        "archive_secondary_mass_kg_m2_source": f"подбор!H{family_row}" if family_row else None,
        "archive_purlin_mass_kg_source": "вывод!E24",
        "archive_secondary_mass_kg_source": None,
        "archive_frame_count_source": None,
        "archive_frame_total_kg_source": None,
    })
    row = _find_city_row(value_wb, city)
    if row:
        climate = value_wb["снегветер"]
        record["archive"].update({
            "canonical_snow_region": _clean(climate[f"F{row}"].value),
            "canonical_snow_load_kpa": _numeric(climate[f"G{row}"].value),
            "canonical_wind_region": _clean(climate[f"H{row}"].value),
            "canonical_wind_load_kpa": _numeric(climate[f"I{row}"].value),
        })
        record["archive"]["snow_load_kpa"] = record["archive"]["canonical_snow_load_kpa"]
        record["archive"]["wind_load_kpa"] = record["archive"]["canonical_wind_load_kpa"]
        record["provenance"].update({
            "archive_snow_load_kpa_source": f"снегветер!G{row}",
            "archive_wind_load_kpa_source": f"снегветер!I{row}",
            "archive_canonical_snow_region_source": f"снегветер!F{row}",
            "archive_canonical_wind_region_source": f"снегветер!H{row}",
        })
    else:
        record["quality_flags"].append("UNKNOWN_CITY")

    for field in ("beam_profile", "column_profile", "purlin_profile"):
        profile = parse_profile(record["archive"].get(field))
        record["archive"][f"{field}_parsed"] = profile
        if profile["parse_error"]:
            record["quality_flags"].append("PROFILE_PARSE_ERROR")
    record["archive_raw_outputs"] = {k: output_values.get(k) for k in output_values}
    return classify(record)
