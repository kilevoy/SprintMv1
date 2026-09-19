"""Extract the proven 9–21 m legacy connection lookup model.

The script is intentionally read-only with respect to the source workbook.  It
walks the two legacy candidate paths (ROW14 and ROW15), follows the nested
INDEX/MATCH formulas to their static branch tables, and writes a versioned JSON
dataset.  Cached values of selector/output cells are not used as universal
project data.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "Таблица по подбору сечений теплых ангаров пролетами 9м, 12м, 15м, 18м, 21м, 24м версия 1,5.xlsx"
DEFAULT_OUTPUT = ROOT / "core1" / "data" / "legacy_connections" / "connection_lookup_rows.json"
DESIGN_FAMILIES = (9, 12, 15, 18, 21)
EXPECTED_SOURCE_SHA256 = "0271b96c6fe725d3e50ad7891401958322a5ae97866bb0bf8cb190bc4bbeff3f"


@dataclass(frozen=True)
class CandidateLayout:
    candidate: str
    height_key_column: str
    metric_columns: dict[str, str]


LAYOUTS = (
    CandidateLayout(
        candidate="ROW14",
        height_key_column="HJ",
        metric_columns={
            "ridgeBeamBoltQuantity": "HN",
            "fittingsWeightKg": "HR",
            "ridgeBeamBoltPattern": "HS",
            "eaveBeamBoltPattern": "HT",
            "supportColumnBoltPattern": "HU",
            "eaveColumnBoltPattern": "HV",
        },
    ),
    CandidateLayout(
        candidate="ROW15",
        height_key_column="QZ",
        metric_columns={
            "ridgeBeamBoltQuantity": "RD",
            "fittingsWeightKg": "RH",
            "ridgeBeamBoltPattern": "RI",
            "eaveBeamBoltPattern": "RJ",
            "supportColumnBoltPattern": "RK",
            "eaveColumnBoltPattern": "RL",
        },
    ),
)

FACTOR_HEIGHT_ROWS = {
    1.0: (7, 8, 9),
    0.8: (15, 16, 17),
}

DIRECT_REFERENCE = re.compile(r"^=\$?([A-Z]+)\$?(\d+)$")
INDEX_MATCH = re.compile(
    r"^=INDEX\((\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+),"
    r"MATCH\([^,]+,(\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+),0\)\)$",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_range(reference: str) -> str:
    return reference.replace("$", "")


def direct_reference(formula: Any, address: str) -> str:
    if not isinstance(formula, str):
        raise ValueError(f"{address}: expected direct-reference formula, got {formula!r}")
    match = DIRECT_REFERENCE.fullmatch(formula)
    if not match:
        raise ValueError(f"{address}: unsupported direct-reference formula {formula!r}")
    return f"{match.group(1)}{match.group(2)}"


def index_match_ranges(formula: Any, address: str) -> tuple[str, str]:
    if not isinstance(formula, str):
        raise ValueError(f"{address}: expected INDEX/MATCH formula, got {formula!r}")
    match = INDEX_MATCH.fullmatch(formula.replace(" ", ""))
    if not match:
        raise ValueError(f"{address}: unsupported INDEX/MATCH formula {formula!r}")
    return clean_range(match.group(1)), clean_range(match.group(2))


def range_cells(sheet: Any, reference: str) -> list[Any]:
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    if min_col != max_col:
        raise ValueError(f"Expected one-column range, got {reference}")
    return [sheet.cell(row=row, column=min_col) for row in range(min_row, max_row + 1)]


def branch_map(formula_sheet: Any, value_sheet: Any, source_address: str) -> tuple[list[str], dict[str, Any], dict[str, str]]:
    value_range, key_range = index_match_ranges(formula_sheet[source_address].value, source_address)
    key_cells = range_cells(value_sheet, key_range)
    value_cells = range_cells(value_sheet, value_range)
    if len(key_cells) != len(value_cells):
        raise ValueError(f"{source_address}: key/value lengths differ: {key_range} vs {value_range}")
    keys: list[str] = []
    values: dict[str, Any] = {}
    sources: dict[str, str] = {}
    for key_cell, value_cell in zip(key_cells, value_cells, strict=True):
        key = key_cell.value
        if key is None:
            continue
        key_text = str(key)
        if key_text in values:
            raise ValueError(f"{source_address}: duplicate branch key {key_text!r}")
        keys.append(key_text)
        values[key_text] = value_cell.value
        sources[key_text] = value_cell.coordinate
    return keys, values, sources


def extract(workbook_path: Path) -> dict[str, Any]:
    source_sha = sha256(workbook_path)
    # Normal mode is intentional: the workbook is wide and the extractor uses
    # sparse random cell access.  read_only mode would restart XML iteration for
    # each access and turns this bounded extraction into an O(n²) operation.
    formula_book = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
    value_book = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    rows: list[dict[str, Any]] = []
    anomalies: list[dict[str, Any]] = []

    for family in DESIGN_FAMILIES:
        sheet_name = f"{family}м"
        formula_sheet = formula_book[sheet_name]
        value_sheet = value_book[sheet_name]
        for layout in LAYOUTS:
            for factor, height_rows in FACTOR_HEIGHT_ROWS.items():
                for height_row in height_rows:
                    height = value_sheet[f"{layout.height_key_column}{height_row}"].value
                    metric_maps: dict[str, tuple[list[str], dict[str, Any], dict[str, str], str]] = {}
                    for metric, column in layout.metric_columns.items():
                        intermediate_address = f"{column}{height_row}"
                        source_address = direct_reference(formula_sheet[intermediate_address].value, intermediate_address)
                        keys, values, sources = branch_map(formula_sheet, value_sheet, source_address)
                        metric_maps[metric] = (keys, values, sources, source_address)

                    ordered_keys = metric_maps["ridgeBeamBoltQuantity"][0]
                    all_keys = set().union(*(set(metric_map[1]) for metric_map in metric_maps.values()))
                    for branch_key in ordered_keys:
                        missing_metrics = [metric for metric, metric_map in metric_maps.items() if branch_key not in metric_map[1]]
                        if missing_metrics:
                            anomalies.append({
                                "classification": "LEGACY_LOOKUP_NO_MATCH",
                                "designFamily": family,
                                "candidate": layout.candidate,
                                "factor": factor,
                                "heightBandM": height,
                                "branchKey": branch_key,
                                "missingMetrics": missing_metrics,
                            })
                            continue
                        metric_values = {metric: metric_map[1][branch_key] for metric, metric_map in metric_maps.items()}
                        metric_sources = {metric: metric_map[2][branch_key] for metric, metric_map in metric_maps.items()}
                        if not isinstance(metric_values["ridgeBeamBoltQuantity"], (int, float)) or not isinstance(metric_values["fittingsWeightKg"], (int, float)):
                            raise ValueError(f"{sheet_name}/{layout.candidate}/{factor}/{height}/{branch_key}: numeric output expected")
                        non_text_patterns = [name for name in (
                            "ridgeBeamBoltPattern", "eaveBeamBoltPattern", "supportColumnBoltPattern", "eaveColumnBoltPattern"
                        ) if not isinstance(metric_values[name], str)]
                        if non_text_patterns:
                            anomalies.append({
                                "classification": "LEGACY_NON_TEXT_BOLT_PATTERN",
                                "designFamily": family,
                                "candidate": layout.candidate,
                                "factor": factor,
                                "heightBandM": height,
                                "branchKey": branch_key,
                                "affectedMetrics": non_text_patterns,
                                "values": {name: metric_values[name] for name in non_text_patterns},
                            })
                        rows.append({
                            "designFamily": family,
                            "candidate": layout.candidate,
                            "factor": factor,
                            "heightBandM": height,
                            "branchKey": branch_key,
                            **metric_values,
                            "source": {
                                "sheet": sheet_name,
                                "valueCells": metric_sources,
                            },
                        })

                    omitted_keys = sorted(all_keys.difference(ordered_keys))
                    for branch_key in omitted_keys:
                        anomalies.append({
                            "classification": "LEGACY_LOOKUP_NO_MATCH",
                            "designFamily": family,
                            "candidate": layout.candidate,
                            "factor": factor,
                            "heightBandM": height,
                            "branchKey": branch_key,
                            "missingMetrics": ["ridgeBeamBoltQuantity"],
                        })

    shared_equal = 0
    row14 = {(row["designFamily"], row["factor"], row["heightBandM"], row["branchKey"]): row for row in rows if row["candidate"] == "ROW14"}
    row15 = {(row["designFamily"], row["factor"], row["heightBandM"], row["branchKey"]): row for row in rows if row["candidate"] == "ROW15"}
    comparable_fields = (
        "ridgeBeamBoltQuantity", "fittingsWeightKg", "ridgeBeamBoltPattern",
        "eaveBeamBoltPattern", "supportColumnBoltPattern", "eaveColumnBoltPattern",
    )
    for key in row14.keys() & row15.keys():
        if any(row14[key][field] != row15[key][field] for field in comparable_fields):
            raise ValueError(f"ROW14/ROW15 values differ for normalized key {key}")
        shared_equal += 1

    return {
        "schemaVersion": 1,
        "classification": "LEGACY_CONNECTION_MODEL_COMPLETE",
        "supportedDomain": {
            "designFamilies": list(DESIGN_FAMILIES),
            "frameStepMode": "AUTOMATIC_ONLY",
            "climateMode": "CITY_LOOKUP_WITH_PROVEN_LEGACY_CLIMATE",
            "heightBandsM": [3.6, 4.8, 6.0],
            "factors": [1.0, 0.8],
            "candidates": [layout.candidate for layout in LAYOUTS],
            "excluded": ["24m_LEGACY_NA", "MANUAL_FRAME_STEP_UNVERIFIED", "MANUAL_CLIMATE_UNVERIFIED"],
        },
        "source": {
            "workbook": workbook_path.name,
            "sha256": source_sha,
            "expectedSha256": EXPECTED_SOURCE_SHA256,
            "sheets": [f"{family}м" for family in DESIGN_FAMILIES],
            "formulaPath": "candidate → factor → height band → mapped legacy branch key → F/J/K/L/M/N",
        },
        "statistics": {
            "rowCount": len(rows),
            "row14Count": len(row14),
            "row15Count": len(row15),
            "sharedEqualKeys": shared_equal,
            "anomalyCount": len(anomalies),
        },
        "anomalies": anomalies,
        "rows": rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    dataset = extract(args.workbook.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(dataset["statistics"], ensure_ascii=False))


if __name__ == "__main__":
    main()
