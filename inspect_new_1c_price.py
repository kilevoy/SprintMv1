from pathlib import Path
import json
from openpyxl import load_workbook

p = Path(r"E:\SprintMv1\input_price_review\ПРАЙС ПОЛНЫЙ на 23.09.2026.xlsx")
wb = load_workbook(p, read_only=True, data_only=False)
out = []
for ws in wb.worksheets:
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True):
        rows.append(list(row[: min(20, ws.max_column)]))
    out.append({"sheet": ws.title, "max_row": ws.max_row, "max_col": ws.max_column, "rows": rows})
print(json.dumps(out, ensure_ascii=False, indent=2, default=str))
